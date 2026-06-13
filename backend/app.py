from datetime import datetime
from io import StringIO, BytesIO
import csv
import os
import sqlite3

from flask import Flask, jsonify, request, Response, send_file, send_from_directory
from flask_cors import CORS
from openpyxl import Workbook


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(BASE_DIR)
FRONTEND_DIST = os.path.join(PROJECT_DIR, "frontend", "dist")
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE_DIR, "data"))
DB_PATH = os.environ.get("DB_PATH", os.path.join(DATA_DIR, "records.db"))


def create_app():
    app = Flask(__name__, static_folder=FRONTEND_DIST, static_url_path="")
    CORS(app)
    init_db()

    @app.get("/api/health")
    def health():
        return jsonify({"status": "ok"})

    @app.get("/api/lines")
    def get_lines():
        rows = query_all(
            "SELECT id, name, color, created_at FROM lines ORDER BY created_at ASC, id ASC"
        )
        return jsonify([dict(row) for row in rows])

    @app.post("/api/lines")
    def create_line():
        payload = request.get_json(silent=True) or {}
        name = clean_text(payload.get("name"))
        color = clean_color(payload.get("color"))
        if not name:
            return error("折线名称不能为空", 400)

        now = datetime.utcnow().isoformat(timespec="seconds") + "Z"
        with get_db() as conn:
            cursor = conn.execute(
                "INSERT INTO lines (name, color, created_at) VALUES (?, ?, ?)",
                (name, color, now),
            )
            conn.commit()
            line = query_one(
                "SELECT id, name, color, created_at FROM lines WHERE id = ?",
                (cursor.lastrowid,),
                conn,
            )
        return jsonify(dict(line)), 201

    @app.put("/api/lines/<int:line_id>")
    def update_line(line_id):
        payload = request.get_json(silent=True) or {}
        existing = query_one("SELECT id FROM lines WHERE id = ?", (line_id,))
        if not existing:
            return error("折线不存在", 404)

        name = clean_text(payload.get("name"))
        color = clean_color(payload.get("color"))
        if not name:
            return error("折线名称不能为空", 400)

        with get_db() as conn:
            conn.execute(
                "UPDATE lines SET name = ?, color = ? WHERE id = ?",
                (name, color, line_id),
            )
            conn.commit()
            line = query_one(
                "SELECT id, name, color, created_at FROM lines WHERE id = ?",
                (line_id,),
                conn,
            )
        return jsonify(dict(line))

    @app.delete("/api/lines/<int:line_id>")
    def delete_line(line_id):
        existing = query_one("SELECT id FROM lines WHERE id = ?", (line_id,))
        if not existing:
            return error("折线不存在", 404)
        with get_db() as conn:
            conn.execute("DELETE FROM lines WHERE id = ?", (line_id,))
            conn.commit()
        return jsonify({"deleted": True})

    @app.get("/api/data-points")
    def get_data_points():
        line_ids = parse_int_list(request.args.get("line_ids"))
        start_date = request.args.get("start_date")
        end_date = request.args.get("end_date")
        where, params = build_data_point_filters(line_ids, start_date, end_date)
        rows = query_all(
            f"""
            SELECT dp.id, dp.line_id, dp.date, dp.value, l.name AS line_name, l.color AS line_color
            FROM data_points dp
            JOIN lines l ON l.id = dp.line_id
            {where}
            ORDER BY dp.date ASC, dp.id ASC
            """,
            params,
        )
        return jsonify([row_to_data_point(row) for row in rows])

    @app.post("/api/data-points")
    def create_data_point():
        payload = request.get_json(silent=True) or {}
        parsed, response = parse_data_point_payload(payload, require_line=True)
        if response:
            return response

        existing = query_one(
            "SELECT id FROM data_points WHERE line_id = ? AND date = ?",
            (parsed["line_id"], parsed["date"]),
        )
        overwrite = bool(payload.get("overwrite", False))
        if existing and not overwrite:
            return (
                jsonify(
                    {
                        "error": "同一条折线在同一天已有数据点",
                        "code": "duplicate_data_point",
                        "existing_id": existing["id"],
                    }
                ),
                409,
            )

        with get_db() as conn:
            if existing:
                conn.execute(
                    "UPDATE data_points SET value = ? WHERE id = ?",
                    (parsed["value"], existing["id"]),
                )
                point_id = existing["id"]
            else:
                cursor = conn.execute(
                    "INSERT INTO data_points (line_id, date, value) VALUES (?, ?, ?)",
                    (parsed["line_id"], parsed["date"], parsed["value"]),
                )
                point_id = cursor.lastrowid
            conn.commit()
            point = fetch_data_point(point_id, conn)
        return jsonify(row_to_data_point(point)), 200 if existing else 201

    @app.put("/api/data-points/<int:point_id>")
    def update_data_point(point_id):
        payload = request.get_json(silent=True) or {}
        current = query_one("SELECT id, line_id FROM data_points WHERE id = ?", (point_id,))
        if not current:
            return error("数据点不存在", 404)

        parsed, response = parse_data_point_payload(
            {**payload, "line_id": payload.get("line_id", current["line_id"])},
            require_line=True,
        )
        if response:
            return response

        duplicate = query_one(
            """
            SELECT id FROM data_points
            WHERE line_id = ? AND date = ? AND id != ?
            """,
            (parsed["line_id"], parsed["date"], point_id),
        )
        overwrite = bool(payload.get("overwrite", False))
        if duplicate and not overwrite:
            return (
                jsonify(
                    {
                        "error": "该日期已有另一个数据点",
                        "code": "duplicate_data_point",
                        "existing_id": duplicate["id"],
                    }
                ),
                409,
            )

        with get_db() as conn:
            if duplicate:
                conn.execute("DELETE FROM data_points WHERE id = ?", (point_id,))
                conn.execute(
                    "UPDATE data_points SET value = ? WHERE id = ?",
                    (parsed["value"], duplicate["id"]),
                )
                saved_id = duplicate["id"]
            else:
                conn.execute(
                    "UPDATE data_points SET line_id = ?, date = ?, value = ? WHERE id = ?",
                    (parsed["line_id"], parsed["date"], parsed["value"], point_id),
                )
                saved_id = point_id
            conn.commit()
            point = fetch_data_point(saved_id, conn)
        return jsonify(row_to_data_point(point))

    @app.delete("/api/data-points/<int:point_id>")
    def delete_data_point(point_id):
        existing = query_one("SELECT id FROM data_points WHERE id = ?", (point_id,))
        if not existing:
            return error("数据点不存在", 404)
        with get_db() as conn:
            conn.execute("DELETE FROM data_points WHERE id = ?", (point_id,))
            conn.commit()
        return jsonify({"deleted": True})

    @app.get("/api/export/csv")
    def export_csv():
        rows = export_rows_from_request()
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["折线名称", "日期", "数值"])
        for row in rows:
            writer.writerow([row["line_name"], row["date"], row["value"]])

        return Response(
            output.getvalue(),
            mimetype="text/csv; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=data-export.csv"},
        )

    @app.get("/api/export/excel")
    def export_excel():
        rows = export_rows_from_request()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["折线名称", "日期", "数值"])
        for row in rows:
            sheet.append([row["line_name"], row["date"], row["value"]])

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return send_file(
            stream,
            as_attachment=True,
            download_name="data-export.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.get("/")
    def serve_index():
        if os.path.exists(os.path.join(FRONTEND_DIST, "index.html")):
            return send_from_directory(FRONTEND_DIST, "index.html")
        return error("前端静态文件不存在，请先运行 npm run build", 404)

    @app.get("/<path:path>")
    def serve_frontend(path):
        target = os.path.join(FRONTEND_DIST, path)
        if os.path.isfile(target):
            return send_from_directory(FRONTEND_DIST, path)
        if os.path.exists(os.path.join(FRONTEND_DIST, "index.html")):
            return send_from_directory(FRONTEND_DIST, "index.html")
        return error("前端静态文件不存在，请先运行 npm run build", 404)

    return app


def init_db():
    os.makedirs(DATA_DIR, exist_ok=True)
    with get_db() as conn:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS lines (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                color TEXT NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS data_points (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                line_id INTEGER NOT NULL,
                date TEXT NOT NULL,
                value REAL NOT NULL,
                FOREIGN KEY (line_id) REFERENCES lines(id) ON DELETE CASCADE,
                UNIQUE (line_id, date)
            );

            CREATE INDEX IF NOT EXISTS idx_data_points_date ON data_points(date);
            CREATE INDEX IF NOT EXISTS idx_data_points_line_date ON data_points(line_id, date);
            """
        )
        conn.commit()


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def query_all(sql, params=(), conn=None):
    owns_conn = conn is None
    conn = conn or get_db()
    try:
        return conn.execute(sql, params).fetchall()
    finally:
        if owns_conn:
            conn.close()


def query_one(sql, params=(), conn=None):
    owns_conn = conn is None
    conn = conn or get_db()
    try:
        return conn.execute(sql, params).fetchone()
    finally:
        if owns_conn:
            conn.close()


def fetch_data_point(point_id, conn=None):
    return query_one(
        """
        SELECT dp.id, dp.line_id, dp.date, dp.value, l.name AS line_name, l.color AS line_color
        FROM data_points dp
        JOIN lines l ON l.id = dp.line_id
        WHERE dp.id = ?
        """,
        (point_id,),
        conn,
    )


def row_to_data_point(row):
    return {
        "id": row["id"],
        "line_id": row["line_id"],
        "date": row["date"],
        "value": row["value"],
        "line_name": row["line_name"],
        "line_color": row["line_color"],
    }


def export_rows_from_request():
    line_ids = parse_int_list(request.args.get("line_ids"))
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    where, params = build_data_point_filters(line_ids, start_date, end_date)
    return query_all(
        f"""
        SELECT l.name AS line_name, dp.date, dp.value
        FROM data_points dp
        JOIN lines l ON l.id = dp.line_id
        {where}
        ORDER BY dp.date ASC, l.name ASC, dp.id ASC
        """,
        params,
    )


def build_data_point_filters(line_ids, start_date, end_date):
    clauses = []
    params = []
    if line_ids:
        placeholders = ",".join("?" for _ in line_ids)
        clauses.append(f"dp.line_id IN ({placeholders})")
        params.extend(line_ids)
    if start_date:
        if not is_valid_date(start_date):
            return "WHERE 1 = 0", []
        clauses.append("dp.date >= ?")
        params.append(start_date)
    if end_date:
        if not is_valid_date(end_date):
            return "WHERE 1 = 0", []
        clauses.append("dp.date <= ?")
        params.append(end_date)
    return ("WHERE " + " AND ".join(clauses), params) if clauses else ("", params)


def parse_data_point_payload(payload, require_line=False):
    line_id = payload.get("line_id")
    value = payload.get("value")
    point_date = payload.get("date")

    if require_line:
        try:
            line_id = int(line_id)
        except (TypeError, ValueError):
            return None, error("请选择折线", 400)
        if not query_one("SELECT id FROM lines WHERE id = ?", (line_id,)):
            return None, error("折线不存在", 404)

    if not is_valid_date(point_date):
        return None, error("日期格式必须为 YYYY-MM-DD", 400)
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None, error("数值必须为数字", 400)

    return {"line_id": line_id, "date": point_date, "value": value}, None


def parse_int_list(value):
    if not value:
        return []
    ids = []
    for item in value.split(","):
        item = item.strip()
        if not item:
            continue
        try:
            ids.append(int(item))
        except ValueError:
            continue
    return ids


def clean_text(value):
    return str(value or "").strip()


def clean_color(value):
    value = str(value or "").strip()
    if len(value) == 7 and value.startswith("#"):
        return value
    return "#3b82f6"


def is_valid_date(value):
    if not value:
        return False
    try:
        datetime.strptime(value, "%Y-%m-%d")
        return True
    except ValueError:
        return False


def error(message, status):
    return jsonify({"error": message}), status


app = create_app()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5001, debug=True)
